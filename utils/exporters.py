"""
Exporters — write assessment outputs to Word (.docx) and Excel (.xlsx).
"""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any


class WordExporter:
    """Export assessment summary to a .docx file."""

    def export(self, state: dict[str, Any], output_path: Path) -> None:
        try:
            from docx import Document
            from docx.shared import Pt, RGBColor
        except ImportError:
            raise ImportError("python-docx required: pip install python-docx")

        doc = Document()
        doc.add_heading("MedTech Opportunity Assessment", 0)

        # Scope
        scope = state.get("scope_agent", {})
        if scope:
            doc.add_heading("Project Scope", level=1)
            doc.add_paragraph(f"Geography: {', '.join(scope.get('geography', []))}")
            doc.add_paragraph(f"Technology: {scope.get('technology_category', '')}")
            doc.add_paragraph(f"Customer Segments: {', '.join(scope.get('customer_segments', []))}")

        # Market dynamics
        dynamics = state.get("market_dynamics_agent", {})
        if dynamics:
            doc.add_heading("Market Dynamics", level=1)
            doc.add_heading("Tailwinds", level=2)
            for t in dynamics.get("tailwinds", []):
                doc.add_paragraph(t, style="List Bullet")
            doc.add_heading("Headwinds", level=2)
            for h in dynamics.get("headwinds", []):
                doc.add_paragraph(h, style="List Bullet")

        # Opportunity matrix
        matrix = state.get("scoring_agent", {})
        if matrix:
            doc.add_heading("Opportunity Matrix", level=1)
            table = doc.add_table(rows=1, cols=4)
            table.style = "Table Grid"
            hdr = table.rows[0].cells
            hdr[0].text = "Opportunity"
            hdr[1].text = "Weighted Score"
            hdr[2].text = "Decision Call"
            hdr[3].text = "Recommended Move"
            for opp in matrix.get("scored_opportunities", []):
                row = table.add_row().cells
                row[0].text = opp.get("name", "")
                row[1].text = str(opp.get("weighted_score", ""))
                row[2].text = opp.get("decision_call", "")
                row[3].text = opp.get("recommended_move", "")

        # GTM
        gtm = state.get("gtm_agent", {})
        if gtm:
            doc.add_heading("GTM Playbook", level=1)
            doc.add_paragraph(f"Recommended Entry Model: {gtm.get('recommended_entry_model', '')}")
            doc.add_paragraph(gtm.get("entry_model_rationale", ""))
            doc.add_heading("Executive Recommendations", level=2)
            for rec in gtm.get("executive_recommendations", []):
                doc.add_paragraph(rec, style="List Bullet")

        doc.save(str(output_path))


class ExcelExporter:
    """Export opportunity matrix and signal data to .xlsx."""

    def export(self, state: dict[str, Any], output_path: Path) -> None:
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill
        except ImportError:
            raise ImportError("openpyxl required: pip install openpyxl")

        wb = openpyxl.Workbook()

        # Sheet 1: Opportunity Matrix
        ws = wb.active
        ws.title = "Opportunity Matrix"
        headers = ["Rank", "Opportunity", "Attractiveness", "Ability to Win",
                   "Strategic Fit", "Speed to Revenue", "Weighted Score", "Decision Call"]
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="43224D")

        scored = state.get("scoring_agent", {}).get("scored_opportunities", [])
        for rank, opp in enumerate(scored, 1):
            ws.append([
                rank,
                opp.get("name", ""),
                opp.get("attractiveness", ""),
                opp.get("ability_to_win", ""),
                opp.get("strategic_fit", ""),
                opp.get("speed_to_revenue", ""),
                opp.get("weighted_score", ""),
                opp.get("decision_call", ""),
            ])

        # Sheet 2: Signals
        ws2 = wb.create_sheet("Signals")
        sig_headers = ["Signal ID", "Statement", "Category", "Commercial Implication",
                       "Confidence", "Relevance Score"]
        for col, h in enumerate(sig_headers, 1):
            cell = ws2.cell(row=1, column=col, value=h)
            cell.font = Font(bold=True)

        signals = state.get("signal_agent", {}).get("signals", [])
        for sig in signals:
            ws2.append([
                sig.get("signal_id", ""),
                sig.get("statement", ""),
                sig.get("category", ""),
                sig.get("commercial_implication", ""),
                sig.get("confidence", ""),
                sig.get("relevance_score", ""),
            ])

        wb.save(str(output_path))
